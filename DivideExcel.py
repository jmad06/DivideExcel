#!/usr/bin/env python3
"""
Divide un archivo Excel en un CSV por cada hoja.
Uso: arrastrar uno o varios .xlsx / .xlsm sobre DivideExcel.bat
"""

import csv
import json
import re
import sys
from datetime import date, datetime, time
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Falta la libreria openpyxl. Instalala con:  pip install openpyxl")
    input("Pulsa Enter para salir...")
    sys.exit(1)

# ---- Configuracion (ver config.json junto a este script) ----
CONFIG_PATH = Path(__file__).resolve().parent / "config.json"

CONFIG_POR_DEFECTO = {
    "delimitador": ";",        # ";" para abrir directo en Excel (España). "," para pandas / Power BI.
    "codificacion": "utf-8-sig",   # utf-8-sig conserva tildes y ñ al abrir en Excel
    "saltar_hojas_vacias": True,
    "extensiones": [".xlsx", ".xlsm", ".xltx", ".xltm"],
}


def cargar_config():
    """Lee config.json junto al script. Si falta o esta incompleto, usa/crea los valores por defecto."""
    if not CONFIG_PATH.exists():
        try:
            CONFIG_PATH.write_text(
                json.dumps(CONFIG_POR_DEFECTO, indent=4, ensure_ascii=False), encoding="utf-8"
            )
        except OSError as e:
            print(f"  [!] No se pudo crear config.json, se usan valores por defecto: {e}")
        return dict(CONFIG_POR_DEFECTO)

    try:
        datos = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError) as e:
        print(f"  [!] config.json invalido, se usan valores por defecto: {e}")
        return dict(CONFIG_POR_DEFECTO)

    config = dict(CONFIG_POR_DEFECTO)
    config.update({k: v for k, v in datos.items() if k in CONFIG_POR_DEFECTO})
    return config


CONFIG = cargar_config()
DELIMITADOR = CONFIG["delimitador"]
CODIFICACION = CONFIG["codificacion"]
SALTAR_HOJAS_VACIAS = CONFIG["saltar_hojas_vacias"]
EXTENSIONES = set(CONFIG["extensiones"])
# -----------------------


def nombre_seguro(nombre):
    """Convierte el nombre de hoja en un nombre de archivo valido en Windows."""
    limpio = re.sub(r'[\\/:*?"<>|]', "_", nombre).strip().rstrip(".")
    return limpio or "hoja_sin_nombre"


def formatear(valor):
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "TRUE" if valor else "FALSE"
    if isinstance(valor, datetime):
        if (valor.hour, valor.minute, valor.second) == (0, 0, 0):
            return valor.strftime("%Y-%m-%d")
        return valor.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(valor, date):
        return valor.strftime("%Y-%m-%d")
    if isinstance(valor, time):
        return valor.strftime("%H:%M:%S")
    return valor


def filas_utiles(hoja):
    """Devuelve las filas sin filas ni columnas vacias al final."""
    filas = []
    for fila in hoja.iter_rows(values_only=True):
        valores = [formatear(v) for v in fila]
        while valores and valores[-1] == "":
            valores.pop()
        filas.append(valores)
    while filas and not filas[-1]:
        filas.pop()
    return filas


def contar_formulas_sin_resolver(hoja_formulas, hoja_valores):
    """Cuenta celdas con formula cuyo valor calculado no esta disponible (None)."""
    total = 0
    for fila_f, fila_v in zip(
        hoja_formulas.iter_rows(values_only=True), hoja_valores.iter_rows(values_only=True)
    ):
        for vf, vv in zip(fila_f, fila_v):
            if isinstance(vf, str) and vf.startswith("=") and vv is None:
                total += 1
    return total


def procesar(ruta):
    if not ruta.exists():
        print(f"  [!] No existe: {ruta}")
        return
    if ruta.suffix.lower() not in EXTENSIONES:
        print(f"  [!] Extension no soportada, se ignora: {ruta.name}")
        return

    print(f"\n> {ruta.name}")
    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except Exception as e:
        print(f"  [!] No se pudo abrir: {e}")
        return

    try:
        libro_formulas = load_workbook(ruta, read_only=True, data_only=False)
    except Exception:
        libro_formulas = None

    destino = ruta.parent / f"{ruta.stem}_csv"
    destino.mkdir(exist_ok=True)

    usados = set()
    generados = 0
    omitidas = 0

    hojas_formulas = libro_formulas.worksheets if libro_formulas else []

    for hoja, hoja_f in zip(libro.worksheets, hojas_formulas or [None] * len(libro.worksheets)):
        try:
            filas = filas_utiles(hoja)
            if not filas and SALTAR_HOJAS_VACIAS:
                print(f"  - {hoja.title}: vacia, se omite")
                continue

            if hoja_f is not None:
                sin_resolver = contar_formulas_sin_resolver(hoja_f, hoja)
                if sin_resolver:
                    print(
                        f"  [!] {hoja.title}: {sin_resolver} formula(s) sin valor calculado "
                        "(se exportan vacias; abre y guarda el Excel para recalcularlas)"
                    )

            base = nombre_seguro(hoja.title)
            nombre = base
            n = 2
            while nombre.lower() in usados:
                nombre = f"{base}_{n}"
                n += 1
            usados.add(nombre.lower())

            ancho = max((len(f) for f in filas), default=0)
            salida = destino / f"{nombre}.csv"

            with open(salida, "w", newline="", encoding=CODIFICACION) as f:
                escritor = csv.writer(f, delimiter=DELIMITADOR)
                for fila in filas:
                    escritor.writerow(fila + [""] * (ancho - len(fila)))

            print(f"  - {hoja.title}: {len(filas)} filas x {ancho} col -> {salida.name}")
            generados += 1
        except Exception as e:
            omitidas += 1
            print(f"  [!] {hoja.title}: error al procesar, se omite ({e})")
            continue

    libro.close()
    if libro_formulas:
        libro_formulas.close()
    resumen = f"  {generados} CSV en: {destino}"
    if omitidas:
        resumen += f"  ({omitidas} hoja(s) omitida(s) por error)"
    print(resumen)


def excels_en_carpeta(carpeta):
    """Busca archivos Excel soportados dentro de una carpeta (recursivo), ignorando los de bloqueo (~$)."""
    return sorted(
        p
        for p in carpeta.rglob("*")
        if p.is_file() and p.suffix.lower() in EXTENSIONES and not p.name.startswith("~$")
    )


def main():
    argumentos = [Path(a) for a in sys.argv[1:]]
    if not argumentos:
        print("Arrastra uno o varios archivos Excel (o carpetas) sobre el .bat.")
        print("\nHecho.")
        input("Pulsa Enter para cerrar...")
        return

    rutas = []
    for arg in argumentos:
        if arg.is_dir():
            encontrados = excels_en_carpeta(arg)
            if not encontrados:
                print(f"[!] {arg}: no contiene archivos Excel soportados")
            rutas.extend(encontrados)
        else:
            rutas.append(arg)

    for ruta in rutas:
        procesar(ruta)
    print("\nHecho.")
    input("Pulsa Enter para cerrar...")


if __name__ == "__main__":
    main()